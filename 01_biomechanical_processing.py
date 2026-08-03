#!/usr/bin/env python3
"""
Tendon tensile-test processing for the Col1a2 null/oim collagen homotrimer study.

Reads the compiled wide-format Force/Displacement/Size CSVs and the sample
metadata, derives per-tendon biomechanical parameters (preconditioning,
hysteresis, stress relaxation over a 60 s hold, pull-to-failure), and writes a
per-sample summary as CSV and XLSX with force-time and force-extension plots
embedded.

Usage:
    python 01_biomechanical_processing.py
    python 01_biomechanical_processing.py --raw-dir test_data --results-dir results/test
"""

from __future__ import annotations

import argparse
import logging
import os
import re
import sys
from collections import Counter
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.signal import savgol_filter

import matplotlib
matplotlib.use("Agg")  
import matplotlib.pyplot as plt

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

# np.trapezoid (NumPy >= 2.0) with a fallback to the older np.trapz.
_trapezoid = getattr(np, "trapezoid", None) or np.trapz

log = logging.getLogger("biomech")

# ---------------------------------------------------------------------
# Config
# These constants are matched to those outlined in the methods section of the paper.
# They can be adjusted for sensitivity analysis or similar.
# ---------------------------------------------------------------------

# Smoothing
PRECON_SAVGOL_WINDOW = 101
PRECON_SAVGOL_POLY = 3
HOLD_SAVGOL_WINDOW = 101
HOLD_SAVGOL_POLY = 3

# QC filter thresholds 
SNR_THRESHOLD = 3
RESIDUAL_THRESHOLD = 0.08
RANGE_THRESHOLD = 0.01

# SNR settings. Noise is estimated on the quiescent 1-Preload phase.
SNR_SAVGOL_WINDOW = 101       # smoothing window for the peak (signal) estimate
SNR_SAVGOL_POLY = 3
SNR_BASELINE_MAX_POINTS = 50  # preload samples nearest loading, used for noise SD
FORCE_RESOLUTION_N = 0.005    # load-cell quantisation step (N); floors the noise estimate

# Stress relaxation
HOLD_DURATION_S = 60.0        # hold duration
PEAK_WINDOW_S = 2.0           # window from hold start for peak stress
S60_HALF_WINDOW_S = 1.0       # half-width of the window S60 is averaged over (2s window centred on 60s)

# Maximum modulus
MODULUS_STRAIN_WINDOW = 0.01  # strain window width
MODULUS_PEAK_LOW_FRAC = 0.05  # lower bound, fraction of peak stress
MODULUS_PEAK_HIGH_FRAC = 0.60  # upper bound, fraction of peak stress
MODULUS_MIN_POINTS_IN_WINDOW = 3

# Instrument label strings used to segment each test.
SETNAME_PRECONDITIONING = "5x pre-conditioning"
SETNAME_STRESS_RELAX = "Stress-relax"
SETNAME_FAILURE = "fail"
CYCLE_PRELOAD = "1-Preload"
CYCLE_HOLD = "1-Hold"
CYCLE_1_STRETCH = "1-Stretch"
CYCLE_1_RECOVER = "1-Recover"

# Default file names
DEFAULT_FORCE_FILE = "compiled_Force_N_phase_140526.1.csv"
DEFAULT_DISP_FILE = "compiled_Displacement_mm_phase_140526.1.csv"
DEFAULT_SIZE_FILE = "compiled_Size_mm_phase_140526.1.csv"
DEFAULT_META_FILE = "compiled_metadata_with_filenames.csv"

# Plot sizing
PLOT_FIGSIZE = (4, 3)
PLOT_DPI = 120
XL_IMAGE_WIDTH_PX = 300
XL_IMAGE_HEIGHT_PX = 225
XL_ROW_HEIGHT_PT = 170

# Columns that should not be coerced to numeric on output
NON_NUMERIC_RESULT_COLS = {
    "File name", "Date", "Sample ID", "Mouse ID", "Sex", "Age", "Genotype",
    "Force-Time plot", "Force-Extension plot",
}


# ---------------------------------------------------------------------
# Plotting helpers
# ---------------------------------------------------------------------

def save_force_time_plot(df_full: pd.DataFrame, outfile_png: str) -> str | None:
    """Save a Force vs Time plot for the entire test (preconditioning + hold + failure)."""
    if df_full is None or df_full.empty:
        return None
    try:
        times = df_full["Time_S"].to_numpy(float)
        force = df_full["Force_N"].to_numpy(float)
    except (KeyError, ValueError, TypeError):
        return None
    if len(times) < 5:
        return None

    fig, ax = plt.subplots(figsize=PLOT_FIGSIZE, dpi=PLOT_DPI)
    ax.plot(times, force, linewidth=1)
    ax.set_xlabel("Time (s)")
    ax.set_ylabel("Force (N)")
    ax.set_title("Force vs Time (entire test)")
    fig.tight_layout()
    fig.savefig(outfile_png, dpi=PLOT_DPI)
    plt.close(fig)
    return outfile_png


def save_force_extension_plot(df_fail: pd.DataFrame, outfile_png: str,
                              disp_mod: float | None = None,
                              force_mod: float | None = None) -> str | None:
    """Force vs Extension for the failure segment, optionally marking the modulus point."""
    if df_fail is None or df_fail.empty:
        return None
    try:
        disp = df_fail["Displacement_mm"].to_numpy(float)
        force = df_fail["Force_N"].to_numpy(float)
    except (KeyError, ValueError, TypeError):
        return None
    if len(disp) < 5:
        return None

    fig, ax = plt.subplots(figsize=PLOT_FIGSIZE, dpi=PLOT_DPI)
    ax.plot(disp, force, linewidth=1)
    if disp_mod is not None and force_mod is not None and np.isfinite(disp_mod) and np.isfinite(force_mod):
        ax.scatter([disp_mod], [force_mod], s=50, marker="x", color="red",
                   label="max modulus", zorder=5)
        ax.legend(fontsize=7, loc="best")
    ax.set_xlabel("Extension (mm)")
    ax.set_ylabel("Force (N)")
    ax.set_title("Force vs Extension (failure segment)")
    fig.tight_layout()
    fig.savefig(outfile_png, dpi=PLOT_DPI)
    plt.close(fig)
    return outfile_png


# ---------------------------------------------------------------------
# Metadata
# ---------------------------------------------------------------------

def _normalise_date_id(v) -> str:
    """Return Date_ID as a six-character, zero-padded string."""
    if pd.isna(v):
        return ""
    if isinstance(v, (int, np.integer)):
        return f"{int(v):06d}"
    s = str(v).strip()
    if s.isdigit():
        return s.zfill(6)
    return s

def _normalise_replicate(v) -> str:
    """Return Replicate as a bare integer string ('1', not '1.0')."""
    if pd.isna(v):
        return ""
    if isinstance(v, float) and float(v).is_integer():
        return str(int(v))
    if isinstance(v, (int, np.integer)):
        return str(int(v))
    return str(v).strip()

def load_metadata(meta_path: str | Path) -> pd.DataFrame:
    """
    Load the sample metadata CSV. 'Age' is left categorical ('8wks', '18wks', '52wks').
    """
    meta_path = Path(meta_path)
    if not meta_path.exists():
        raise FileNotFoundError(f"Metadata file '{meta_path}' not found.")

    df = pd.read_csv(meta_path, encoding="utf-8-sig")

    mouse_matches = [c for c in df.columns if c.strip().lower() in ("mouse id", "mouse_id", "mouseid")]
    if mouse_matches:
        if mouse_matches[0] != "Mouse ID":
            df = df.rename(columns={mouse_matches[0]: "Mouse ID"})
    else:
        if df.shape[1] <= 2:
            raise ValueError(
                "Metadata has no column named 'Mouse ID' and too few columns to fall back "
                f"on the original positional rename. Columns found: {list(df.columns)}"
            )
        log.warning(
            "No 'Mouse ID' column found by name; falling back to the original script's "
            "positional rename of column index 2 (%r). Verify this is correct.",
            df.columns[2],
        )
        df = df.rename(columns={df.columns[2]: "Mouse ID"})

    df = df.apply(lambda s: s.map(lambda x: x.strip() if isinstance(x, str) else x))

    rename_map = {}
    for col in df.columns:
        cl = col.strip().lower()
        if cl == "date_id":
            rename_map[col] = "Date_ID"
        elif cl == "sample_id":
            rename_map[col] = "Sample_ID"
        elif cl == "replicate":
            rename_map[col] = "Replicate"
        elif cl == "average_diameter (um)":
            rename_map[col] = "Average_diameter (um)"
        elif cl == "c.s.a (um squared)":
            rename_map[col] = "C.s.a (um squared)"
        elif cl == "c.s.a (mm squared)":
            rename_map[col] = "C.s.a (mm squared)"
    df = df.rename(columns=rename_map)

    required = ["Date_ID", "Sample_ID", "Replicate", "C.s.a (mm squared)"]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"Metadata is missing required column(s): {missing}")

    # Keep Date_ID zero-padded as a string, or the metadata join will fail.
    df["Date_ID"] = df["Date_ID"].apply(_normalise_date_id)
    df["Sample_ID"] = df["Sample_ID"].astype(str).str.strip()
    df["Replicate"] = df["Replicate"].apply(_normalise_replicate)

    for c in ["Average_diameter (um)", "C.s.a (um squared)", "C.s.a (mm squared)"]:
        if c in df.columns:
            df[c] = pd.to_numeric(df[c], errors="coerce")

    return df


# ---------------------------------------------------------------------
# Savitzky-Golay smoothing 
# ---------------------------------------------------------------------

def apply_savgol_safe(y, window: int, poly: int) -> np.ndarray:
    """Savitzky-Golay filter that degrades on short segments."""
    y = np.asarray(y, dtype=float)
    n = len(y)
    if n < 5:
        return y.copy()
    w = min(window, n)
    if w % 2 == 0:
        w -= 1
    if w < 5:
        return y.copy()
    p = min(poly, w - 1)
    return savgol_filter(y, window_length=w, polyorder=p)


def savgol_deriv_safe(y, window: int, poly: int, delta: float) -> np.ndarray:
    """First derivative via Savitzky-Golay, with the same short-segment safety as apply_savgol_safe."""
    y = np.asarray(y, dtype=float)
    n = len(y)
    if n < 5 or not np.isfinite(delta) or delta == 0:
        return np.full(n, np.nan)
    w = min(window, n)
    if w % 2 == 0:
        w -= 1
    if w < 5:
        return np.full(n, np.nan)
    p = min(poly, w - 1)
    if p < 1:
        return np.full(n, np.nan)
    return savgol_filter(y, window_length=w, polyorder=p, deriv=1, delta=delta)


# ---------------------------------------------------------------------
# QC filters
# ---------------------------------------------------------------------

def compute_snr(force, cycle=None, window: int = SNR_SAVGOL_WINDOW,
                poly: int = SNR_SAVGOL_POLY,
                baseline_max: int = SNR_BASELINE_MAX_POINTS) -> float:
    """
    Peak-to-baseline signal-to-noise ratio.

        signal = peak of the smoothed force, above the preload baseline level
        noise  = SD of the raw force over the quiescent 1-Preload phase

    `cycle` is the segment's cycle label; when absent, the leading `baseline_max`
    samples are used as the baseline. Returns NaN when no usable baseline or
    positive signal is found (the segment then cannot pass the SNR filter).
    """
    force = np.asarray(force, dtype=float)

    # ---- locate the signal-free baseline ----
    baseline_vals = None
    if cycle is not None:
        pre_mask = pd.Series(cycle).astype(str).str.contains(
            CYCLE_PRELOAD, case=False, na=False).to_numpy()
        if pre_mask.any():
            b = force[pre_mask]
            b = b[np.isfinite(b)]
            if len(b) >= 2:
                # samples nearest load onset best represent the noise
                baseline_vals = b[-baseline_max:] if len(b) > baseline_max else b
    if baseline_vals is None:
        ff = force[np.isfinite(force)]
        if len(ff) >= 2:
            baseline_vals = ff[:baseline_max]
    if baseline_vals is None or len(baseline_vals) < 2:
        return np.nan

    noise = float(np.std(baseline_vals))
    # Floor at the quantisation noise so a flat preload doesn't give a near-zero SD.
    noise = max(noise, FORCE_RESOLUTION_N / np.sqrt(12.0))
    if not np.isfinite(noise) or noise <= 0:
        return np.nan

    # ---- signal: peak of the smoothed trace above the preload level ----
    baseline_level = float(np.mean(baseline_vals))
    y = force[np.isfinite(force)] - baseline_level
    if len(y) < 5:
        return np.nan
    f_smooth = apply_savgol_safe(y, window, poly)
    signal = float(np.max(f_smooth))
    if not np.isfinite(signal) or signal <= 0:
        return np.nan
    return signal / noise


def compute_residual_ratio(y) -> float:
    """SD of the residual about a Savitzky-Golay fit, as a fraction of signal range."""
    y = np.asarray(y, dtype=float)
    if len(y) < 10:
        return np.nan
    y_s = apply_savgol_safe(y, PRECON_SAVGOL_WINDOW, PRECON_SAVGOL_POLY)
    if np.array_equal(y_s, y):
        return np.nan
    noise = np.std(y - y_s)
    signal = np.max(y) - np.min(y)
    if signal <= 0 or not np.isfinite(signal):
        return np.nan
    return noise / signal


def compute_signal_range(y) -> float:
    y = np.asarray(y, dtype=float)
    if len(y) < 10:
        return np.nan
    return np.max(y) - np.min(y)


def segment_qc(y, cycle, label: str, gate_snr: bool = True) -> dict:
    """
    Compute the three QC metrics (SNR, residual ratio, range) and their pass flags for one segment.

    `gate_snr` controls whether SNR contributes to the overall pass/fail. It is
    only meaningful for the failure curve, where 1-Preload phase is a true baseline.
    SNR is still computed and reported for every segment as a diagnostic, but does not gate preconditioning or hold.

    """
    snr = compute_snr(y, cycle)
    resid = compute_residual_ratio(y)
    rng = compute_signal_range(y)

    snr_pass = bool(snr >= SNR_THRESHOLD) if np.isfinite(snr) else False
    resid_pass = bool(resid <= RESIDUAL_THRESHOLD) if np.isfinite(resid) else False
    rng_pass = bool(rng >= RANGE_THRESHOLD) if np.isfinite(rng) else False
    overall = (snr_pass or not gate_snr) and resid_pass and rng_pass

    if not overall:
        log.info("  FILTER FAILED: %s (snr=%.3g%s resid=%.3g range=%.3g)",
                 label, snr, "" if gate_snr else " [not gated]", resid, rng)

    return {
        "snr": snr, "resid": resid, "range": rng,
        "snr_pass": snr_pass, "resid_pass": resid_pass, "range_pass": rng_pass,
        "snr_gated": gate_snr,
        "pass": overall,
    }


def empty_qc() -> dict:
    return {"snr": np.nan, "resid": np.nan, "range": np.nan,
            "snr_pass": False, "resid_pass": False, "range_pass": False,
            "snr_gated": True, "pass": False}


# ---------------------------------------------------------------------
# Sliding-window modulus
# ---------------------------------------------------------------------

# Returned when no valid modulus window is found.
_NULL_MODULUS = {
    "modulus": np.nan, "strain": np.nan, "stress": np.nan, "global_idx": None,
    "intercept": np.nan, "r2": np.nan, "n_points": 0,
}


def compute_sliding_window_modulus(fail_df: pd.DataFrame, CSA_true: float,
                                   sample_length: float,
                                   window_size_strain: float = MODULUS_STRAIN_WINDOW) -> dict:
    """
    Maximum tangent modulus from a strain window slid along the pre-peak
    stress-strain curve, restricted to 5-60% of peak stress.

    Returns a dict with the modulus (slope), strain/stress at the modulus point,
    its global index, and goodness-of-fit for the winning window (intercept, R2,number of points).
    """
    force = fail_df["Force_N"].to_numpy(float)
    disp = fail_df["Displacement_mm"].to_numpy(float)

    if not np.isfinite(CSA_true) or CSA_true == 0:
        return dict(_NULL_MODULUS)
    if not np.isfinite(sample_length) or sample_length == 0:
        return dict(_NULL_MODULUS)

    load_corr = force - force[0]
    disp_corr = disp - disp[0]

    strain_raw = disp_corr / sample_length
    stress_raw = load_corr / CSA_true

    mask = np.isfinite(strain_raw) & np.isfinite(stress_raw)
    strain = strain_raw[mask]
    stress = stress_raw[mask]
    idx_all = fail_df.index.to_numpy()[mask]

    if len(strain) < 6:
        return dict(_NULL_MODULUS)

    # Pre-peak region only
    peak_idx = int(np.nanargmax(stress))
    if peak_idx <= 2:
        return dict(_NULL_MODULUS)
    strain, stress, idx_all = strain[:peak_idx], stress[:peak_idx], idx_all[:peak_idx]
    if len(strain) < 6:
        return dict(_NULL_MODULUS)

    # Peak-stress window filter (5-60% of peak stress)
    peak_stress = float(np.nanmax(stress))
    low_s = MODULUS_PEAK_LOW_FRAC * peak_stress
    high_s = MODULUS_PEAK_HIGH_FRAC * peak_stress
    mask_zone = (stress >= low_s) & (stress <= high_s)
    if np.sum(mask_zone) < 6:
        return dict(_NULL_MODULUS)
    strain, stress, idx_all = strain[mask_zone], stress[mask_zone], idx_all[mask_zone]

    # Sliding linear fit
    n = len(strain)
    slopes = np.full(n, np.nan)
    half_win = window_size_strain / 2.0

    for i in range(n):
        x0 = strain[i]
        mask_win = (strain >= x0 - half_win) & (strain <= x0 + half_win)
        xs = strain[mask_win]
        ys = stress[mask_win]
        if len(xs) < MODULUS_MIN_POINTS_IN_WINDOW:
            continue
        try:
            m, _c = np.polyfit(xs, ys, 1)
            slopes[i] = m
        except (np.linalg.LinAlgError, ValueError):
            continue

    if not np.isfinite(slopes).any():
        return dict(_NULL_MODULUS)

    idx_local = int(np.nanargmax(slopes))
    global_idx = int(idx_all[idx_local])

    # Refit the winning window to report its goodness-of-fit.
    x0 = strain[idx_local]
    mask_win = (strain >= x0 - half_win) & (strain <= x0 + half_win)
    xs = strain[mask_win]
    ys = stress[mask_win]
    slope, intercept = np.polyfit(xs, ys, 1)
    y_hat = slope * xs + intercept
    ss_res = float(np.sum((ys - y_hat) ** 2))
    ss_tot = float(np.sum((ys - np.mean(ys)) ** 2))
    r2 = 1.0 - ss_res / ss_tot if ss_tot > 0 else np.nan

    return {
        "modulus": float(slope),
        "strain": float(strain[idx_local]),
        "stress": float(stress[idx_local]),
        "global_idx": global_idx,
        "intercept": float(intercept),
        "r2": float(r2),
        "n_points": int(len(xs)),
    }


# ---------------------------------------------------------------------
# Segment analyses
# ---------------------------------------------------------------------

def analyze_preconditioning(df: pd.DataFrame, CSA_true: float) -> dict:
    """Preconditioning QC, sample length, and cycle-1 hysteresis."""
    pre = df[df["SetName"].astype(str).str.contains(
        SETNAME_PRECONDITIONING, case=False, na=False)].copy()
    if pre.empty:
        return {"_ok": False, "_reason": "no preconditioning segment"}

    qc = segment_qc(pre["Force_N"], pre["Cycle"], "pre-conditioning", gate_snr=False)

    sample_length = float(pre["Size_mm"].iloc[0])
    if not np.isfinite(sample_length) or sample_length == 0:
        return {"_ok": False, "_reason": "preconditioning Size_mm is zero/NaN"}

    # Baseline: minimum force across the whole preconditioning block.
    min_force = float(pre["Force_N"].min())
    pre["Load_corr"] = pre["Force_N"] - min_force
    pre["Disp_corr"] = pre["Displacement_mm"] - pre["Displacement_mm"].iloc[0]

    pre["Load_smooth"] = apply_savgol_safe(pre["Load_corr"], PRECON_SAVGOL_WINDOW, PRECON_SAVGOL_POLY)
    pre["Disp_smooth"] = apply_savgol_safe(pre["Disp_corr"], PRECON_SAVGOL_WINDOW, PRECON_SAVGOL_POLY)

    hyst_energy = np.nan
    hyst_pct = np.nan

    cycles = pre["Cycle"].astype(str)
    c1_stretch = cycles.str.contains(CYCLE_1_STRETCH, case=False, na=False)
    c1_recover = cycles.str.contains(CYCLE_1_RECOVER, case=False, na=False)

    if c1_stretch.any() and c1_recover.any() and np.isfinite(CSA_true) and CSA_true != 0:
        xs_load = pre.loc[c1_stretch, "Disp_smooth"].to_numpy(float) / sample_length
        ys_load = pre.loc[c1_stretch, "Load_smooth"].to_numpy(float) / CSA_true
        xs_unload = pre.loc[c1_recover, "Disp_smooth"].to_numpy(float) / sample_length
        ys_unload = pre.loc[c1_recover, "Load_smooth"].to_numpy(float) / CSA_true

        strain_origin = xs_load[0]
        xs_load = xs_load - strain_origin
        xs_unload = xs_unload - strain_origin

        A_load = float(_trapezoid(ys_load, xs_load))
        # Recovery runs in decreasing strain, so its area is negative; take abs().
        A_unload = float(_trapezoid(ys_unload, xs_unload))

        hyst_energy = A_load - abs(A_unload)
        hyst_pct = hyst_energy / A_load * 100 if A_load != 0 else np.nan

    return {
        "_ok": True,
        "sample_length": sample_length,
        "min_force": min_force,
        "hyst_energy": hyst_energy,
        "hyst_pct": hyst_pct,
        "qc": qc,
    }


def analyze_hold(df: pd.DataFrame, CSA_true: float) -> dict:
    """Stress relaxation over the 60 s hold."""
    relax_block = df[df["SetName"].astype(str).str.contains(
        SETNAME_STRESS_RELAX, case=False, na=False)].copy()

    out = {
        "stress_relax_60": np.nan,
        "stress_rate_60": np.nan,
        "s60_clamped": False,
        "qc": empty_qc(),
    }
    if relax_block.empty:
        log.info("  ! No 'Stress-relax' segment found.")
        return out

    out["qc"] = segment_qc(relax_block["Force_N"], relax_block["Cycle"], "stress-relax", gate_snr=False)

    cycles = relax_block["Cycle"].astype(str)
    preload = relax_block[cycles.str.contains(CYCLE_PRELOAD, case=False, na=False)]
    baseline_force = float(preload["Force_N"].mean()) if not preload.empty else 0.0

    hold = relax_block[cycles.str.contains(CYCLE_HOLD, case=False, na=False)].copy()
    if hold.empty or not np.isfinite(CSA_true) or CSA_true == 0:
        return out

    hold["Force_corr"] = hold["Force_N"] - baseline_force
    hold["Stress_raw"] = hold["Force_corr"] / CSA_true
    hold["Stress_smooth"] = apply_savgol_safe(hold["Stress_raw"], HOLD_SAVGOL_WINDOW, HOLD_SAVGOL_POLY)

    times = hold["Time_S"].to_numpy(float)
    stress_sm = hold["Stress_smooth"].to_numpy(float)
    if len(times) < 2:
        return out

    t0 = times[0]
    t60 = t0 + HOLD_DURATION_S
    if times[-1] < t60:
        log.info("  ! Hold segment shorter than %.0f s; stress relaxation not computed.", HOLD_DURATION_S)
        return out

    # Peak stress: maximum of the smoothed trace within the first PEAK_WINDOW_S of the hold
    mask_0 = times <= t0 + PEAK_WINDOW_S
    s_peak = float(np.max(stress_sm[mask_0])) if np.any(mask_0) else float(stress_sm[0])

    # S60: mean of the smoothed trace over a 2 s window centred on 60 s.
    mask_60 = (times >= t60 - S60_HALF_WINDOW_S) & (times <= t60 + S60_HALF_WINDOW_S)
    if np.any(mask_60):
        s60_avg = float(np.mean(stress_sm[mask_60]))
    else:
        s60_avg = float(np.interp(t60, times, stress_sm))

    if s60_avg < 0:
        s60_avg = 0.0
        out["s60_clamped"] = True

    out["stress_relax_60"] = (s_peak - s60_avg) / s_peak * 100 if s_peak != 0 else np.nan
    out["stress_rate_60"] = (s60_avg - s_peak) / HOLD_DURATION_S
    return out


def analyze_failure(df: pd.DataFrame, CSA_true: float, sample_length: float) -> dict:
    """Pull-to-failure properties and maximum modulus."""
    fail = df[df["SetName"].astype(str).str.lower().str.contains(SETNAME_FAILURE, na=False)].copy()
    if fail.empty:
        return {"_ok": False, "_reason": "no failure segment (SetName did not contain 'fail')"}

    qc = segment_qc(fail["Force_N"], fail["Cycle"], "failure")

    fail["Load_corr"] = fail["Force_N"] - fail["Force_N"].iloc[0]
    fail["Disp_corr"] = fail["Displacement_mm"] - fail["Displacement_mm"].iloc[0]
    fail["Strain"] = fail["Disp_corr"] / sample_length
    fail["Stress_MPa"] = fail["Load_corr"] / CSA_true if (np.isfinite(CSA_true) and CSA_true != 0) else np.nan

    fidx = fail["Load_corr"].idxmax()
    failure_force = float(fail.loc[fidx, "Load_corr"])
    failure_ext = float(fail.loc[fidx, "Disp_corr"])
    failure_strain = float(fail.loc[fidx, "Strain"])
    failure_stress = (float(fail.loc[fidx, "Load_corr"]) / CSA_true
                      if (np.isfinite(CSA_true) and CSA_true != 0) else np.nan)

    max_mod = strain_at_mod = stress_at_mod = np.nan
    disp_at_mod = force_at_mod = time_at_mod = np.nan
    mod_intercept = mod_r2 = np.nan
    mod_n_points = 0

    mod = compute_sliding_window_modulus(fail, CSA_true, sample_length)
    global_idx = mod["global_idx"]
    if global_idx is not None:
        max_mod, strain_at_mod, stress_at_mod = mod["modulus"], mod["strain"], mod["stress"]
        mod_intercept, mod_r2, mod_n_points = mod["intercept"], mod["r2"], mod["n_points"]
        force_at_mod = float(fail["Force_N"].loc[global_idx])
        time_at_mod = float(fail["Time_S"].loc[global_idx])
        disp_at_mod = float(fail["Displacement_mm"].loc[global_idx])

    return {
        "_ok": True,
        "fail_df": fail,
        "failure_force": failure_force,
        "failure_ext": failure_ext,
        "failure_strain": failure_strain,
        "failure_stress": failure_stress,
        "max_mod": max_mod,
        "strain_at_mod": strain_at_mod,
        "stress_at_mod": stress_at_mod,
        "disp_at_mod": disp_at_mod,
        "force_at_mod": force_at_mod,
        "time_at_mod": time_at_mod,
        "mod_intercept": mod_intercept,
        "mod_r2": mod_r2,
        "mod_n_points": mod_n_points,
        "qc": qc,
    }


# ---------------------------------------------------------------------
# Analysis of a single sample
# ---------------------------------------------------------------------

def analyze_dataframe(df: pd.DataFrame, meta: pd.Series, plot_dir: Path,
                      make_plots: bool = True) -> dict | None:
    sample_name = f"{meta['Date_ID']}_Sample_{meta['Sample_ID']}{meta['Replicate']}"
    log.info("Processing %s Sample %s%s ...", meta["Date_ID"], meta["Sample_ID"], meta["Replicate"])

    required = ["SetName", "Cycle", "Time_S", "Size_mm", "Displacement_mm", "Force_N"]
    for c in required:
        if c not in df.columns:
            raise ValueError(f"Column {c} missing.")

    avg_diam = float(meta["Average_diameter (um)"])
    CSA_raw = float(meta["C.s.a (um squared)"])
    CSA_true = float(meta["C.s.a (mm squared)"])
  
    if not np.isfinite(CSA_true) or CSA_true == 0:
        log.warning("  ! %s: cross-sectional area is %r; stress endpoints will be NaN.",
                    sample_name, CSA_true)

    pre_res = analyze_preconditioning(df, CSA_true)
    if not pre_res["_ok"]:
        log.info("  ! %s: %s", sample_name, pre_res["_reason"])
        return None
    sample_length = pre_res["sample_length"]

    hold_res = analyze_hold(df, CSA_true)

    fail_res = analyze_failure(df, CSA_true, sample_length)
    if not fail_res["_ok"]:
        log.info("  ! %s: %s", sample_name, fail_res["_reason"])
        return None

    ft_path = plot_dir / f"{sample_name}_force_time.png"
    fe_path = plot_dir / f"{sample_name}_force_extension.png"
    if make_plots:
        plot_dir.mkdir(parents=True, exist_ok=True)
        save_force_time_plot(df, str(ft_path))
        # Mark the modulus point for a per-sample visual QC check.
        save_force_extension_plot(fail_res["fail_df"], str(fe_path),
                                  disp_mod=fail_res["disp_at_mod"],
                                  force_mod=fail_res["force_at_mod"])

    result = {
        "File name": sample_name,
        "Date": meta["Date"],
        "Sample ID": meta["Sample_ID"],
        "Replicate number": int(meta["Replicate"]),
        "Mouse ID": meta.get("Mouse ID", ""),
        "Sex": meta["Sex"],
        "Age": meta["Age"],
        "Genotype": meta["Genotype"],

        "Stress-relaxation 60s": hold_res["stress_relax_60"],
        "Rate of change of stress 60s": hold_res["stress_rate_60"],
        "S60 clamped": hold_res["s60_clamped"],

        "Hysteresis strain energy per unit volume": pre_res["hyst_energy"],
        "Hysteresis %": pre_res["hyst_pct"],

        "Average diameter (um)": avg_diam,
        "C.s.a (um squared)": CSA_raw,
        "C.s.a (mm squared)": CSA_true,
        "Sample length (mm)": sample_length,

        "Failure force (N) - corrected": fail_res["failure_force"],
        "Failure stress (MPa) - corrected": fail_res["failure_stress"],
        "Failure strain (%) - corrected": fail_res["failure_strain"] * 100,
        "Failure extension (mm) - corrected": fail_res["failure_ext"],

        "Max modulus (sliding window)": fail_res["max_mod"],
        "Stress at max modulus (sliding window)": fail_res["stress_at_mod"],
        "Strain at max modulus (sliding window)": fail_res["strain_at_mod"],
        "Time at max modulus (s)": fail_res["time_at_mod"],
        "Force at max modulus (N)": fail_res["force_at_mod"],
        "Displacement at max modulus (mm)": fail_res["disp_at_mod"],
        "Max modulus fit R2": fail_res["mod_r2"],
        "Max modulus fit intercept (MPa)": fail_res["mod_intercept"],
        "Max modulus fit n points": fail_res["mod_n_points"],

        "SNR preconditioning": pre_res["qc"]["snr"],
        "SNR hold": hold_res["qc"]["snr"],
        "SNR failure": fail_res["qc"]["snr"],
        "SNR preconditioning pass": pre_res["qc"]["snr_pass"],
        "SNR hold pass": hold_res["qc"]["snr_pass"],
        "SNR failure pass": fail_res["qc"]["snr_pass"],

        "Residual preconditioning": pre_res["qc"]["resid"],
        "Residual hold": hold_res["qc"]["resid"],
        "Residual failure": fail_res["qc"]["resid"],
        "Residual preconditioning pass": pre_res["qc"]["resid_pass"],
        "Residual hold pass": hold_res["qc"]["resid_pass"],
        "Residual failure pass": fail_res["qc"]["resid_pass"],

        "Range preconditioning": pre_res["qc"]["range"],
        "Range hold": hold_res["qc"]["range"],
        "Range failure": fail_res["qc"]["range"],
        "Range preconditioning pass": pre_res["qc"]["range_pass"],
        "Range hold pass": hold_res["qc"]["range_pass"],
        "Range failure pass": fail_res["qc"]["range_pass"],

        # Overall per-segment QC verdicts. Reported only; exclusion happens downstream.
        "Preconditioning pass": pre_res["qc"]["pass"],
        "Hold pass": hold_res["qc"]["pass"],
        "Failure pass": fail_res["qc"]["pass"],

        "Force-Time plot": str(ft_path),
        "Force-Extension plot": str(fe_path),
    }

    for k, v in result.items():
        if isinstance(v, np.generic):
            result[k] = v.item()
    return result


# ---------------------------------------------------------------------
# Column matching
# ---------------------------------------------------------------------
# Column suffixes in compiled data are truncated as they were originally seperate worksheets in Excel (capped at 31 chars). 
# Examples:
#   "210330 MRC Sample A1Data_SetNam"   (SetName -> SetNam)
#   "210330 MRC Sample A10Data_SetNa"   (SetName -> SetNa)
#   "210330 MRC Sample A1Data_Force_"   (Force_N -> Force_)
# Matching therefore has to accept any non-empty PREFIX of the canonical suffix.

KNOWN_SUFFIXES = ("SetName", "Cycle", "Force_N", "Displacement_mm", "Size_mm")
KNOWN_SUFFIXES = ("SetName", "Cycle", "Force_N", "Displacement_mm", "Size_mm")

# Below this length a truncated suffix is no longer unique (e.g. "S" -> SetName or Size_mm).
MIN_SUFFIX_CHARS = 3

# Columns in the compiled files that are not per-sample data.
NON_SAMPLE_COLUMNS = {"Time_S"}


def _suffix_matches(remainder: str, expected: str) -> bool:
    """True if `remainder` is a non-empty (possibly truncated) prefix of `expected`."""
    if not remainder:
        return False
    return expected.lower().startswith(remainder.lower())


def find_col(df: pd.DataFrame, base: str, expected_suffix: str) -> str | None:
    """Find the column for `base` carrying `expected_suffix`, tolerating truncation (see KNOWN_SUFFIXES)."""
    prefix = f"{base}_"
    matches = []
    for c in df.columns:
        s = str(c)
        if not s.startswith(prefix):
            continue
        remainder = s[len(prefix):]
        if not _suffix_matches(remainder, expected_suffix):
            continue
        # If the remainder fits several canonical suffixes, flag rather than guess.
        rivals = [x for x in KNOWN_SUFFIXES
                  if x != expected_suffix and _suffix_matches(remainder, x)]
        if rivals:
            log.warning(
                "Column %r truncated to ambiguous suffix %r (could be %s or %s); skipping. "
                "Shorten the sample names before export.",
                s, remainder, expected_suffix, "/".join(rivals),
            )
            continue
        matches.append(s)

    if not matches:
        return None
    if len(matches) > 1:
        log.warning("Ambiguous column match for base %r / %r: %s", base, expected_suffix, matches)
    return matches[0]


def sample_bases(force_df: pd.DataFrame) -> list[str]:
    """Sample base names, taken from each sample's SetName column."""
    bases = []
    for col in force_df.columns:
        s = str(col)
        if s in NON_SAMPLE_COLUMNS:
            continue
        i = s.rfind("_")
        if i <= 0:
            continue
        remainder = s[i + 1:]
        if len(remainder) < MIN_SUFFIX_CHARS:
            # e.g. a base long enough to clip 'SetName' down to 'Se' or 'S'.
            if _suffix_matches(remainder, "SetName"):
                log.warning(
                    "Column %r has a suffix truncated past recognition (%r). The sample "
                    "name is too long for the 31-character export limit; this sample "
                    "cannot be identified and will be skipped.", s, remainder,
                )
            continue
        if _suffix_matches(remainder, "SetName"):
            bases.append(s[:i])
    return bases


# ---------------------------------------------------------------------
# Metadata lookup
# ---------------------------------------------------------------------

# Fallback: pull (Date_ID, Sample_ID, Replicate) from a column base name,
# discarding anything after the replicate digits (e.g. 'Sample C3.2Data' -> (C, 3)).
BASE_NAME_RE = re.compile(r"(\d{6}).*Sample\s*([A-Z])\s*([0-9]+)")


def lookup_metadata(metadata: pd.DataFrame, base: str):
    """
    Find the metadata row for a sample, returning (rows, how).

    Prefers an exact match on the FileName column, falling back to the name
    regex. Returns (None, how) if the base cannot be resolved at all.
    """
    if "FileName" in metadata.columns:
        target = f"{base}.csv"
        rows = metadata[metadata["FileName"].astype(str).str.strip() == target]
        if not rows.empty:
            return rows, "FileName"

    m = BASE_NAME_RE.search(base)
    if not m:
        return None, "FileName + name regex"
    date_id, sample_id, rep = m.group(1), m.group(2), str(int(m.group(3)))
    rows = metadata[
        (metadata["Date_ID"].astype(str) == date_id)
        & (metadata["Sample_ID"].astype(str) == sample_id)
        & (metadata["Replicate"].astype(str) == rep)
    ]
    return rows, f"name regex (Date_ID={date_id}, Sample_ID={sample_id}, Replicate={rep})"


# ---------------------------------------------------------------------
# Output cleaning
# ---------------------------------------------------------------------

def clean_results_frame(df: pd.DataFrame) -> pd.DataFrame:
    """Tidy the results table, coercing text columns to numeric where safe."""
    df = df.copy()
    for col in df.columns:
        if col in NON_NUMERIC_RESULT_COLS:
            continue
        if pd.api.types.is_bool_dtype(df[col]):
            continue
        if pd.api.types.is_numeric_dtype(df[col]):
            continue
        cleaned = (
            df[col].astype(str).str.strip()
            .str.replace("\u2212", "-", regex=False)   # minus sign
            .str.replace("\u2013", "-", regex=False)   # en dash
            .replace({"nan": np.nan, "None": np.nan, "": np.nan})
        )
        converted = pd.to_numeric(cleaned, errors="coerce")
        # Only go numeric if no real value was lost in the conversion.
        if converted.notna().sum() == cleaned.notna().sum():
            df[col] = converted
        else:
            df[col] = cleaned
    return df


def embed_plots(xlsx_path: Path, df: pd.DataFrame) -> None:
    """Embed the force-time and force-extension PNGs into the XLSX summary."""
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    headers = list(df.columns)

    for header in ("Force-Time plot", "Force-Extension plot"):
        if header not in headers:
            continue
        col_idx = headers.index(header) + 1
        for row_idx in range(2, len(df) + 2):
            plot_path = ws.cell(row=row_idx, column=col_idx).value
            if not plot_path or not os.path.exists(str(plot_path)):
                continue
            img = XLImage(str(plot_path))
            img.width = XL_IMAGE_WIDTH_PX
            img.height = XL_IMAGE_HEIGHT_PX
            ws.add_image(img, f"{get_column_letter(col_idx)}{row_idx}")
            ws.row_dimensions[row_idx].height = XL_ROW_HEIGHT_PT

    wb.save(xlsx_path)


# ---------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------

def parse_args(argv=None) -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Process compiled tendon tensile-test data into a per-sample summary.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    p.add_argument("--raw-dir", type=Path, default=Path("raw_data"),
                   help="Directory containing the compiled input CSVs.")
    p.add_argument("--results-dir", type=Path, default=Path("results"),
                   help="Directory for results_summary.csv / .xlsx.")
    p.add_argument("--plots-dir", type=Path, default=None,
                   help="Directory for per-sample PNGs (default: <results-dir>/plots).")
    p.add_argument("--force-file", default=DEFAULT_FORCE_FILE)
    p.add_argument("--disp-file", default=DEFAULT_DISP_FILE)
    p.add_argument("--size-file", default=DEFAULT_SIZE_FILE)
    p.add_argument("--meta-file", default=DEFAULT_META_FILE)
    p.add_argument("--prefix", default="results_summary",
                   help="Basename for the output files.")
    p.add_argument("--no-plots", action="store_true",
                   help="Skip PNG generation and XLSX embedding (for faster testing/performance).")
    p.add_argument("-v", "--verbose", action="store_true")
    return p.parse_args(argv)


def main(argv=None) -> int:
    args = parse_args(argv)
    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(levelname)s: %(message)s",
        stream=sys.stdout,
    )

    plots_dir = args.plots_dir or (args.results_dir / "plots")
    args.results_dir.mkdir(parents=True, exist_ok=True)

    metadata = load_metadata(args.raw_dir / args.meta_file)
    force_df = pd.read_csv(args.raw_dir / args.force_file, low_memory=False, encoding="utf-8-sig")
    disp_df = pd.read_csv(args.raw_dir / args.disp_file, low_memory=False, encoding="utf-8-sig")
    size_df = pd.read_csv(args.raw_dir / args.size_file, low_memory=False, encoding="utf-8-sig")

    if "Time_S" not in force_df.columns:
        raise ValueError(f"'Time_S' column not found in {args.force_file}.")
    time = pd.to_numeric(force_df["Time_S"], errors="coerce")

    results = []
    skips = Counter()
    bases = sample_bases(force_df)
    log.info("Found %d sample columns in %s", len(bases), args.force_file)

    for base in bases:
        set_col = find_col(force_df, base, "SetName")
        cycle_col = find_col(force_df, base, "Cycle")
        force_col = find_col(force_df, base, "Force_N")
        disp_col = find_col(disp_df, base, "Displacement_mm")
        size_col = find_col(size_df, base, "Size_mm")

        missing = [n for n, c in [("SetName", set_col), ("Cycle", cycle_col),
                                  ("Force_N", force_col), ("Displacement_mm", disp_col),
                                  ("Size_mm", size_col)] if c is None]
        if missing:
            log.warning("SKIP %s: missing column(s) %s", base, missing)
            skips["missing_columns"] += 1
            continue

        meta_rows, how = lookup_metadata(metadata, base)
        if meta_rows is None:
            log.warning("SKIP %s: could not parse Date_ID / Sample_ID / Replicate from the name, "
                        "and no FileName match", base)
            skips["unparseable_name"] += 1
            continue
        if meta_rows.empty:
            log.warning("SKIP %s: no metadata row found (tried %s)", base, how)
            skips["no_metadata"] += 1
            continue
        if len(meta_rows) > 1:
            log.warning("%s: %d metadata rows matched via %s; using the first.",
                        base, len(meta_rows), how)
        meta = meta_rows.iloc[0]

        df_sample = pd.DataFrame({
            "Time_S": time,
            "SetName": force_df[set_col],
            "Cycle": force_df[cycle_col],
            "Force_N": pd.to_numeric(force_df[force_col], errors="coerce"),
            "Displacement_mm": pd.to_numeric(disp_df[disp_col], errors="coerce"),
            "Size_mm": pd.to_numeric(size_df[size_col], errors="coerce"),
        })

        try:
            r = analyze_dataframe(df_sample, meta, plot_dir=plots_dir,
                                  make_plots=not args.no_plots)
        except Exception:
            log.exception("SKIP %s: unhandled error during analysis", base)
            skips["analysis_error"] += 1
            continue

        if r is None:
            skips["missing_segment"] += 1
            continue

        clean_base = re.sub(r"Data$", "", base)
        r["File name"] = clean_base + "Data.csv"
        results.append(r)

    if not results:
        log.error("No valid samples.")
        return 1

    df = clean_results_frame(pd.DataFrame(results))

    log.info("Sample columns found : %d", len(bases))
    for reason, n in sorted(skips.items()):
        log.info("Skipped (%s): %d", reason, n)
    log.info("Samples analysed     : %d", len(df))
    for seg in ("Preconditioning pass", "Hold pass", "Failure pass"):
        if seg in df.columns:
            log.info("%-22s: %d / %d pass QC", seg, int(df[seg].sum()), len(df))

    csv_path = args.results_dir / f"{args.prefix}.csv"
    df.to_csv(csv_path, index=False, float_format="%.10g")
    log.info("Wrote %s", csv_path)

    xlsx_path = args.results_dir / f"{args.prefix}.xlsx"
    df.to_excel(xlsx_path, index=False)
    if not args.no_plots:
        embed_plots(xlsx_path, df)
        log.info("Wrote %s (with embedded plots)", xlsx_path)
    else:
        log.info("Wrote %s", xlsx_path)

    return 0

if __name__ == "__main__":
    sys.exit(main())
